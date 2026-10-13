#!/usr/bin/env python3
"""Build a client-ready website audit report (.docx) from a content JSON file.

Usage:
    python3 build_report.py content.json "Output Report.docx" [--check-only]

The JSON schema is documented in SKILL.md. The script validates the content,
prints warnings about jargon / length / structure, and renders a styled Word
document with: cover page, executive summary, what's working, key
opportunities, prioritized recommendations, and recommended next steps.
"""

import argparse
import datetime as _dt
import json
import re
import sys

try:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Inches, Cm, Emu
except ImportError:  # pragma: no cover
    sys.exit("python-docx is required: pip install python-docx")


# --------------------------------------------------------------------------
# Validation
# --------------------------------------------------------------------------

LEVELS = ("High", "Medium", "Low")
TIMINGS = ("Quick win", "Next", "Later")
TIMING_HEADINGS = {
    "Quick win": "Quick wins (next 30 days)",
    "Next": "Next (1–3 months)",
    "Later": "Later (3+ months)",
}

# Terms that usually signal untranslated audit-speak. Matched case-insensitively
# on word boundaries. Not every hit is wrong, but each deserves a second look.
JARGON = [
    r"LCP", r"CLS", r"FID", r"INP", r"TTFB", r"FCP", r"Core Web Vitals",
    r"canonical\w*", r"hreflang", r"robots\.txt", r"noindex", r"nofollow",
    r"schema(\.org)?", r"structured data", r"JSON-LD", r"sitemap\.xml",
    r"30[127]", r"404s?", r"5\d\d errors?", r"render[- ]blocking", r"minif\w+",
    r"lazy[- ]load\w*", r"DOM", r"CSS", r"JavaScript", r"JS", r"HTML", r"H1s?",
    r"H2s?", r"alt (text|attributes?|tags?)", r"meta (tags?|descriptions?)",
    r"title tags?", r"crawl\w*", r"indexab\w+", r"SERPs?", r"CTR", r"DA",
    r"backlinks?", r"WCAG", r"ARIA", r"viewport", r"WebP", r"CDN", r"gzip",
    r"Lighthouse", r"PageSpeed Insights", r"Screaming Frog", r"Ahrefs",
    r"SEMrush", r"GTM", r"GA4", r"UTM",
]
JARGON_RE = re.compile(r"\b(" + "|".join(JARGON) + r")\b", re.IGNORECASE)
URL_RE = re.compile(r"(https?://|www\.)\S+|`[^`]+`|</?[a-z]+[^>]*>")


def _words(text):
    return len(str(text).split())


def _as_paragraphs(value):
    if value is None:
        return []
    if isinstance(value, list):
        return [str(v) for v in value if str(v).strip()]
    return [p for p in re.split(r"\n\s*\n", str(value)) if p.strip()]


def validate(c):
    """Return (errors, warnings)."""
    errors, warnings = [], []

    for key in ("client_name", "executive_summary", "whats_working",
                "opportunities", "recommendations", "next_steps"):
        if not c.get(key):
            errors.append(f"Missing required field: '{key}'")
    if errors:
        return errors, warnings

    es = c["executive_summary"]
    if not es.get("overview"):
        errors.append("executive_summary.overview is required")
    if _words(" ".join(_as_paragraphs(es.get("overview")))) > 140:
        warnings.append("Executive summary overview is over 140 words; aim for 3-5 sentences.")
    kt = es.get("key_takeaways") or []
    if not 2 <= len(kt) <= 4:
        warnings.append(f"executive_summary.key_takeaways has {len(kt)} items; 3 is ideal.")
    if len(es.get("at_a_glance") or []) > 4:
        warnings.append("at_a_glance has more than 4 items; only the first 4 will be shown.")

    ww = c["whats_working"]
    if not 2 <= len(ww) <= 6:
        warnings.append(f"whats_working has {len(ww)} items; 3-5 is ideal.")

    opps = c["opportunities"]
    if not 2 <= len(opps) <= 7:
        warnings.append(f"opportunities has {len(opps)} items; 3-6 is ideal. Merge or cut.")
    for i, o in enumerate(opps, 1):
        for k in ("title", "finding", "why_it_matters", "impact"):
            if not o.get(k):
                errors.append(f"opportunities[{i}] missing '{k}'")
        if o.get("impact") and o["impact"] not in LEVELS:
            errors.append(f"opportunities[{i}].impact must be one of {LEVELS}")
        if _words(o.get("finding", "")) + _words(o.get("why_it_matters", "")) > 120:
            warnings.append(f"Opportunity '{o.get('title')}' is long (>120 words); tighten it.")

    recs = c["recommendations"]
    if not 3 <= len(recs) <= 12:
        warnings.append(f"recommendations has {len(recs)} items; 6-10 is typical.")
    for i, r in enumerate(recs, 1):
        for k in ("action", "detail", "impact", "effort", "timing"):
            if not r.get(k):
                errors.append(f"recommendations[{i}] missing '{k}'")
        if r.get("impact") and r["impact"] not in LEVELS:
            errors.append(f"recommendations[{i}].impact must be one of {LEVELS}")
        if r.get("effort") and r["effort"] not in LEVELS:
            errors.append(f"recommendations[{i}].effort must be one of {LEVELS}")
        if r.get("timing") and r["timing"] not in TIMINGS:
            errors.append(f"recommendations[{i}].timing must be one of {TIMINGS}")
    if recs and not any(r.get("timing") == "Quick win" for r in recs):
        warnings.append("No 'Quick win' recommendations; clients respond well to at least one.")

    ns = c["next_steps"]
    steps = ns.get("steps") if isinstance(ns, dict) else None
    if not steps:
        errors.append("next_steps.steps is required")
    elif not 2 <= len(steps) <= 6:
        warnings.append(f"next_steps has {len(steps)} steps; 3-5 is ideal.")

    # Jargon / leftover technical detail scan over all client-facing text.
    for path, text in _iter_text(c):
        hits = sorted({m.group(0) for m in JARGON_RE.finditer(text)}, key=str.lower)
        if hits:
            warnings.append(f"Possible jargon in {path}: {', '.join(hits)}")
        if URL_RE.search(text):
            warnings.append(f"URL / code / markup in {path}; keep technical detail out of the client report.")

    return errors, warnings


SKIP_KEYS = {"brand_color", "site_url", "date", "impact", "effort", "timing", "page_size"}


def _iter_text(obj, path=""):
    if isinstance(obj, dict):
        for k, v in obj.items():
            if k in SKIP_KEYS:
                continue
            yield from _iter_text(v, f"{path}.{k}" if path else k)
    elif isinstance(obj, list):
        for i, v in enumerate(obj):
            yield from _iter_text(v, f"{path}[{i}]")
    elif isinstance(obj, str):
        yield path, obj


# --------------------------------------------------------------------------
# Low-level docx helpers
# --------------------------------------------------------------------------

def _hex(color):
    color = (color or "").lstrip("#")
    return color.upper() if re.fullmatch(r"[0-9A-Fa-f]{6}", color) else None


def _rgb(hex_color):
    return RGBColor.from_string(hex_color)


def _tint(hex_color, amount):
    """Blend hex_color toward white by `amount` (0..1)."""
    r, g, b = (int(hex_color[i:i + 2], 16) for i in (0, 2, 4))
    mix = lambda v: round(v + (255 - v) * amount)
    return f"{mix(r):02X}{mix(g):02X}{mix(b):02X}"


# Successor tags, so new elements are inserted in schema order (Word rejects
# out-of-order children in some cases).
_TBLPR_AFTER_BORDERS = ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
                        "w:tblCaption", "w:tblDescription", "w:tblPrChange")
_TCPR_AFTER_MAR = ("w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark",
                   "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge", "w:tcPrChange")
_TCPR_AFTER_SHD = ("w:noWrap", "w:tcMar") + _TCPR_AFTER_MAR
_TCPR_AFTER_BORDERS = ("w:shd",) + _TCPR_AFTER_SHD
_PPR_AFTER_PBDR = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
                   "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
                   "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
                   "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
                   "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
                   "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")
_RPR_AFTER_SHD = ("w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang",
                  "w:eastAsianLayout", "w:specVanish", "w:oMath")


def _replace_child(parent, el, successors):
    for old in parent.findall(el.tag):
        parent.remove(old)
    parent.insert_element_before(el, *successors)


def _shd(fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    return shd


def shade_cell(cell, fill):
    _replace_child(cell._tc.get_or_add_tcPr(), _shd(fill), _TCPR_AFTER_SHD)


def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    mar = OxmlElement("w:tcMar")
    for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    _replace_child(cell._tc.get_or_add_tcPr(), mar, _TCPR_AFTER_MAR)


def set_cell_left_bar(cell, color, size=36):
    b = OxmlElement("w:tcBorders")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), "0")
    left.set(qn("w:color"), color)
    b.append(left)
    _replace_child(cell._tc.get_or_add_tcPr(), b, _TCPR_AFTER_BORDERS)


def set_table_borders(table, color="FFFFFF", size=0, inside_h=None):
    """Set all borders; size 0 = none. inside_h=(size,color) overrides horizontal inner lines."""
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        s, col = size, color
        if edge == "insideH" and inside_h:
            s, col = inside_h
        el.set(qn("w:val"), "single" if s else "nil")
        if s:
            el.set(qn("w:sz"), str(s))
            el.set(qn("w:color"), col)
            el.set(qn("w:space"), "0")
        borders.append(el)
    _replace_child(table._tbl.tblPr, borders, _TBLPR_AFTER_BORDERS)


def set_col_widths(table, widths):
    table.autofit = False
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    _replace_child(table._tbl.tblPr, layout, _TBLPR_AFTER_BORDERS[2:])
    widths = [Emu(int(w)) for w in widths]
    for row in table.rows:
        for idx, w in enumerate(widths):
            row.cells[idx].width = w
    for idx, gc in enumerate(table._tbl.tblGrid.findall(qn("w:gridCol"))):
        if idx < len(widths):
            gc.set(qn("w:w"), str(int(widths[idx]) // 635))


def no_split_row(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def repeat_header(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))


def para_border_bottom(paragraph, color, size=8, space=4):
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    _replace_child(paragraph._p.get_or_add_pPr(), pbdr, _PPR_AFTER_PBDR)


def run_shading(run, fill):
    _replace_child(run._r.get_or_add_rPr(), _shd(fill), _RPR_AFTER_SHD)


def add_field(run, instr):
    """Insert a field (e.g. PAGE) into a run."""
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instr
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    val = OxmlElement("w:t")
    val.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, text, sep, val, end):
        run._r.append(el)


def fmt(paragraph, before=0, after=6, line=1.15, keep_next=False, align=None):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = keep_next
    if align is not None:
        paragraph.alignment = align
    return paragraph


def add_run(paragraph, text, size=None, bold=None, italic=None, color=None, font=None):
    run = paragraph.add_run(text)
    if size:
        run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = _rgb(color)
    if font:
        run.font.name = font
    return run


# --------------------------------------------------------------------------
# Report builder
# --------------------------------------------------------------------------

LEVEL_COLORS = {  # (fill, text) for impact badges
    "High": ("FDE2DD", "A12C17"),
    "Medium": ("FFF1CC", "8A5A00"),
    "Low": ("E3F1E6", "2E6B3A"),
}
INK = "1F2933"      # body text
MUTED = "5F6B7A"    # secondary text
RULE = "D9DEE5"     # light lines


class ReportBuilder:
    def __init__(self, content):
        self.c = content
        self.brand = _hex(content.get("brand_color")) or "1F4E79"
        self.brand_tint = _tint(self.brand, 0.90)
        self.brand_mid = _tint(self.brand, 0.55)
        self.doc = Document()
        self._setup()

    # -- setup -------------------------------------------------------------
    def _setup(self):
        doc = self.doc
        sec = doc.sections[0]
        if str(self.c.get("page_size", "")).upper() == "A4":
            sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        else:
            sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        sec.left_margin = sec.right_margin = Inches(1)
        sec.top_margin = Inches(0.9)
        sec.bottom_margin = Inches(0.9)
        sec.header_distance = sec.footer_distance = Inches(0.45)
        self.content_width = sec.page_width - sec.left_margin - sec.right_margin

        normal = doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        normal.font.color.rgb = _rgb(INK)
        rpr = normal.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(attr), "Calibri")
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.15

        for name, size, color in (("Heading 1", 20, self.brand), ("Heading 2", 13.5, INK),
                                  ("Heading 3", 11.5, self.brand)):
            st = doc.styles[name]
            st.font.name = "Calibri"
            st.font.size = Pt(size)
            st.font.bold = True
            st.font.italic = False
            st.font.color.rgb = _rgb(color)
            r = st.element.get_or_add_rPr()
            rf = r.find(qn("w:rFonts"))
            if rf is None:
                rf = OxmlElement("w:rFonts")
                r.append(rf)
            for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                rf.set(qn(attr), "Calibri")
            for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
                if rf.get(qn(attr)) is not None:
                    del rf.attrib[qn(attr)]
            st.paragraph_format.keep_with_next = True

        # python-docx's default template omits the required zoom percent.
        for zoom in doc.settings.element.findall(qn("w:zoom")):
            if zoom.get(qn("w:percent")) is None:
                zoom.set(qn("w:percent"), "100")

        core = doc.core_properties
        core.title = self._title()
        core.subject = f"Website audit for {self.c['client_name']}"
        if self.c.get("prepared_by"):
            core.author = self.c["prepared_by"]

    def _title(self):
        return self.c.get("report_title") or "Website Audit: Findings & Recommendations"

    # -- header / footer ------------------------------------------------------
    def _header_footer(self):
        sec = self.doc.sections[0]
        sec.different_first_page_header_footer = True

        hp = sec.header.paragraphs[0]
        fmt(hp, after=0, align=WD_ALIGN_PARAGRAPH.RIGHT)
        add_run(hp, f"{self.c['client_name']}  ·  {self._title()}", size=8.5, color=MUTED)
        para_border_bottom(hp, RULE, size=4, space=6)

        fp = sec.footer.paragraphs[0]
        fmt(fp, after=0)
        tab_stops = fp.paragraph_format.tab_stops
        # The built-in Footer style has centre/right tabs at Letter positions; clear them.
        for pos in (Inches(3.25), Inches(6.5)):
            tab_stops.add_tab_stop(pos, alignment=WD_TAB_ALIGNMENT.CLEAR)
        tab_stops.add_tab_stop(self.content_width, alignment=WD_TAB_ALIGNMENT.RIGHT)
        left = "Confidential, prepared for " + self.c["client_name"]
        add_run(fp, left, size=8.5, color=MUTED)
        add_run(fp, "\tPage ", size=8.5, color=MUTED)
        r = add_run(fp, "", size=8.5, color=MUTED)
        add_field(r, "PAGE")

    # -- section heading --------------------------------------------------------
    def section_heading(self, number, text):
        p = self.doc.add_paragraph(style="Heading 1")
        fmt(p, before=6, after=10, line=1.0, keep_next=True)
        add_run(p, f"{number:02d}  ", color=self.brand_mid, size=20, bold=False)
        add_run(p, text, bold=False)
        para_border_bottom(p, self.brand, size=12, space=6)
        return p

    def body(self, text, after=6, color=None, italic=None, size=None, keep_next=False):
        p = self.doc.add_paragraph()
        fmt(p, after=after, keep_next=keep_next)
        add_run(p, text, color=color, italic=italic, size=size)
        return p

    def bullet(self, text, bold_lead=None):
        p = self.doc.add_paragraph(style="List Bullet")
        fmt(p, after=4)
        if bold_lead:
            add_run(p, bold_lead, bold=True)
            add_run(p, " " + text)
        else:
            add_run(p, text)
        return p

    def badge(self, paragraph, level, prefix=""):
        fill, txt = LEVEL_COLORS.get(level, (RULE, INK))
        r = add_run(paragraph, f" {prefix}{level} ", size=8.5, bold=True, color=txt)
        run_shading(r, fill)
        return r

    # -- cover -----------------------------------------------------------------
    def cover(self):
        c = self.c
        doc = self.doc

        spacer = doc.add_paragraph()
        fmt(spacer, after=0)
        spacer.paragraph_format.space_before = Pt(60)

        band = doc.add_table(rows=1, cols=1)
        band.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(band)
        cell = band.cell(0, 0)
        shade_cell(cell, self.brand)
        set_cell_margins(cell, top=520, bottom=520, left=480, right=480)
        set_col_widths(band, [self.content_width])

        p = cell.paragraphs[0]
        fmt(p, after=10, line=1.0)
        add_run(p, "WEBSITE AUDIT REPORT", size=10, bold=True, color=_tint(self.brand, 0.65))
        p = cell.add_paragraph()
        fmt(p, after=10, line=1.0)
        add_run(p, self._title(), size=30, bold=True, color="FFFFFF")
        if c.get("subtitle"):
            p = cell.add_paragraph()
            fmt(p, after=0, line=1.1)
            add_run(p, c["subtitle"], size=13, color=_tint(self.brand, 0.80))

        gap = doc.add_paragraph()
        fmt(gap, after=0)
        gap.paragraph_format.space_before = Pt(36)

        meta = [("Prepared for", c["client_name"])]
        if c.get("site_url"):
            meta.append(("Website", c["site_url"]))
        if c.get("prepared_by"):
            meta.append(("Prepared by", c["prepared_by"]))
        meta.append(("Date", c.get("date") or _dt.date.today().strftime("%B %Y")))

        for label, value in meta:
            p = doc.add_paragraph()
            fmt(p, after=8, line=1.0)
            p.paragraph_format.left_indent = Inches(0.1)
            add_run(p, label.upper(), size=8.5, bold=True, color=MUTED)
            p2 = doc.add_paragraph()
            fmt(p2, after=12, line=1.0)
            p2.paragraph_format.left_indent = Inches(0.1)
            add_run(p2, value, size=13, color=INK)
            p.paragraph_format.space_after = Pt(1)

        # Contents strip
        p = doc.add_paragraph()
        fmt(p, before=30, after=4)
        add_run(p, "INSIDE THIS REPORT", size=8.5, bold=True, color=MUTED)
        para_border_bottom(p, RULE, size=4, space=4)
        items = ["Executive summary", "What's working", "Key opportunities",
                 "Prioritized recommendations", "Recommended next steps"]
        for i, item in enumerate(items, 1):
            p = doc.add_paragraph()
            fmt(p, after=2)
            add_run(p, f"{i:02d}   ", size=10, bold=True, color=self.brand)
            add_run(p, item, size=10, color=INK)

        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # -- 1. executive summary ------------------------------------------------------
    def executive_summary(self):
        es = self.c["executive_summary"]
        self.section_heading(1, "Executive summary")
        paras = _as_paragraphs(es.get("overview"))
        for i, text in enumerate(paras):
            p = self.body(text, after=8, size=11.5 if i == 0 else None)

        glance = (es.get("at_a_glance") or [])[:4]
        if glance:
            p = self.doc.add_paragraph()
            fmt(p, before=8, after=4, keep_next=True)
            add_run(p, "AT A GLANCE", size=8.5, bold=True, color=MUTED)
            t = self.doc.add_table(rows=1, cols=len(glance))
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_table_borders(t, color="FFFFFF", size=24)  # white gutters between tiles
            w = self.content_width // len(glance)
            set_col_widths(t, [w] * len(glance))
            for cell, item in zip(t.rows[0].cells, glance):
                shade_cell(cell, self.brand_tint)
                set_cell_margins(cell, top=160, bottom=160, left=160, right=160)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                cp = cell.paragraphs[0]
                fmt(cp, after=2, line=1.0)
                add_run(cp, str(item.get("value", "")), size=20, bold=True, color=self.brand)
                cp2 = cell.add_paragraph()
                fmt(cp2, after=0, line=1.05)
                add_run(cp2, str(item.get("label", "")), size=9, color=MUTED)
            no_split_row(t.rows[0])

        kt = es.get("key_takeaways") or []
        if kt:
            p = self.doc.add_paragraph()
            fmt(p, before=14, after=4, keep_next=True)
            add_run(p, "KEY TAKEAWAYS", size=8.5, bold=True, color=MUTED)
            box = self.doc.add_table(rows=1, cols=1)
            set_table_borders(box)
            set_col_widths(box, [self.content_width])
            cell = box.cell(0, 0)
            shade_cell(cell, "F5F7FA")
            set_cell_margins(cell, top=140, bottom=100, left=220, right=220)
            set_cell_left_bar(cell, self.brand)
            for i, t in enumerate(kt):
                cp = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
                fmt(cp, after=6, line=1.15)
                cp.paragraph_format.left_indent = Inches(0.28)
                cp.paragraph_format.first_line_indent = Inches(-0.28)
                add_run(cp, f"{i + 1}\t", bold=True, color=self.brand)
                cp.paragraph_format.tab_stops.add_tab_stop(Inches(0.28))
                add_run(cp, t)

    # -- 2. what's working -----------------------------------------------------------
    def whats_working(self):
        self.section_heading(2, "What's working")
        intro = self.c.get("whats_working_intro")
        if intro:
            self.body(intro, after=8, keep_next=True)
        items = self.c["whats_working"]
        t = self.doc.add_table(rows=len(items), cols=2)
        set_table_borders(t, inside_h=(4, RULE))
        set_col_widths(t, [Inches(0.45), self.content_width - Inches(0.45)])
        for row, item in zip(t.rows, items):
            no_split_row(row)
            ic, tc = row.cells
            set_cell_margins(ic, top=120, bottom=120, left=0, right=60)
            set_cell_margins(tc, top=120, bottom=120, left=60, right=0)
            ip = ic.paragraphs[0]
            fmt(ip, after=0, line=1.0)
            add_run(ip, "✓", size=15, bold=True, color="2E7D4F")
            p = tc.paragraphs[0]
            fmt(p, after=2, line=1.1)
            add_run(p, item.get("title", ""), bold=True, size=11.5)
            if item.get("detail"):
                p2 = tc.add_paragraph()
                fmt(p2, after=0, line=1.15)
                add_run(p2, item["detail"], color=INK)

    # -- 3. key opportunities ----------------------------------------------------
    def opportunities(self):
        self.section_heading(3, "Key opportunities")
        intro = self.c.get("opportunities_intro") or (
            "These are the areas where focused improvements are most likely to "
            "translate into more enquiries, sales and visibility, ordered by "
            "expected impact.")
        self.body(intro, after=10, color=MUTED, keep_next=True)
        for i, o in enumerate(self.c["opportunities"], 1):
            t = self.doc.add_table(rows=1, cols=1)
            set_table_borders(t, color=RULE, size=6)
            set_col_widths(t, [self.content_width])
            cell = t.cell(0, 0)
            set_cell_margins(cell, top=160, bottom=120, left=220, right=220)
            no_split_row(t.rows[0])

            p = cell.paragraphs[0]
            fmt(p, after=6, line=1.05)
            add_run(p, f"{i}.  ", bold=True, size=13, color=self.brand)
            add_run(p, o["title"], bold=True, size=13)
            add_run(p, "   ")
            self.badge(p, o["impact"], prefix="Impact: ")

            for label, key in (("What we found", "finding"), ("Why it matters", "why_it_matters")):
                lp = cell.add_paragraph()
                fmt(lp, after=1, line=1.0)
                add_run(lp, label.upper(), size=8, bold=True, color=self.brand)
                bp = cell.add_paragraph()
                fmt(bp, after=6, line=1.15)
                add_run(bp, o[key])

            sp = self.doc.add_paragraph()
            fmt(sp, after=0, line=0.6)

    # -- 4. prioritized recommendations --------------------------------------------
    def recommendations(self):
        self.section_heading(4, "Prioritized recommendations")
        recs = self.c["recommendations"]
        order = {k: i for i, k in enumerate(TIMINGS)}
        recs = sorted(recs, key=lambda r: order.get(r.get("timing"), 9))  # stable
        intro = self.c.get("recommendations_intro") or (
            "We've ranked each recommendation by its likely impact on your goals and "
            "the effort involved. Quick wins come first: they deliver visible results "
            "quickly and build momentum for the larger pieces of work.")
        self.body(intro, after=10, keep_next=True)

        # Summary table
        cols = ["#", "Recommendation", "Impact", "Effort", "Timing"]
        widths = [Inches(0.4), None, Inches(0.85), Inches(0.85), Inches(1.0)]
        fixed = sum(w for w in widths if w is not None)
        widths[1] = self.content_width - fixed
        t = self.doc.add_table(rows=1 + len(recs), cols=len(cols))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(t, inside_h=(4, RULE))
        set_col_widths(t, widths)
        hdr = t.rows[0]
        repeat_header(hdr)
        for cell, name in zip(hdr.cells, cols):
            shade_cell(cell, self.brand)
            set_cell_margins(cell, top=90, bottom=90, left=100, right=100)
            p = cell.paragraphs[0]
            fmt(p, after=0, line=1.0)
            add_run(p, name, bold=True, size=9.5, color="FFFFFF")
        for n, (row, r) in enumerate(zip(t.rows[1:], recs), 1):
            no_split_row(row)
            vals = [str(n), r["action"], r["impact"], r["effort"], r["timing"]]
            for j, (cell, v) in enumerate(zip(row.cells, vals)):
                set_cell_margins(cell, top=90, bottom=90, left=100, right=100)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if n % 2 == 0:
                    shade_cell(cell, "F7F9FB")
                p = cell.paragraphs[0]
                fmt(p, after=0, line=1.05)
                if j == 0:
                    add_run(p, v, bold=True, size=10, color=self.brand)
                elif j == 1:
                    add_run(p, v, bold=True, size=10)
                elif j == 2:
                    self.badge(p, v)
                else:
                    add_run(p, v, size=9.5, color=INK)

        # Detail by timing group
        n = 0
        for timing in TIMINGS:
            group = [r for r in recs if r.get("timing") == timing]
            if not group:
                continue
            h = self.doc.add_paragraph(style="Heading 2")
            fmt(h, before=16, after=6, keep_next=True)
            add_run(h, TIMING_HEADINGS[timing], color=self.brand)
            para_border_bottom(h, RULE, size=4, space=3)
            for r in group:
                n += 1
                p = self.doc.add_paragraph()
                fmt(p, before=4, after=2, keep_next=True)
                p.paragraph_format.left_indent = Inches(0.32)
                p.paragraph_format.first_line_indent = Inches(-0.32)
                p.paragraph_format.tab_stops.add_tab_stop(Inches(0.32))
                add_run(p, f"{n}\t", bold=True, color=self.brand)
                add_run(p, r["action"], bold=True)
                d = self.doc.add_paragraph()
                fmt(d, after=4)
                d.paragraph_format.left_indent = Inches(0.32)
                add_run(d, r["detail"])
                m = self.doc.add_paragraph()
                fmt(m, after=6)
                m.paragraph_format.left_indent = Inches(0.32)
                self.badge(m, r["impact"], prefix="Impact: ")
                add_run(m, "   Effort: ", size=8.5, color=MUTED)
                add_run(m, r["effort"], size=8.5, bold=True, color=MUTED)

    # -- 5. next steps ----------------------------------------------------------
    def next_steps(self):
        ns = self.c["next_steps"]
        self.section_heading(5, "Recommended next steps")
        if ns.get("intro"):
            self.body(ns["intro"], after=10, keep_next=True)
        steps = ns["steps"]
        t = self.doc.add_table(rows=len(steps), cols=2)
        set_table_borders(t, color="FFFFFF", size=24)  # white gaps between rows
        set_col_widths(t, [Inches(0.75), self.content_width - Inches(0.75)])
        for i, (row, s) in enumerate(zip(t.rows, steps), 1):
            no_split_row(row)
            nc, tc = row.cells
            shade_cell(nc, self.brand)
            nc.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(nc, top=120, bottom=120, left=60, right=60)
            np_ = nc.paragraphs[0]
            fmt(np_, after=0, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
            add_run(np_, str(i), size=18, bold=True, color="FFFFFF")
            shade_cell(tc, self.brand_tint)
            set_cell_margins(tc, top=140, bottom=140, left=200, right=200)
            tc.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = tc.paragraphs[0]
            fmt(p, after=2, line=1.05)
            add_run(p, s.get("title", ""), bold=True, size=11.5)
            if s.get("timing"):
                add_run(p, f"   ·  {s['timing']}", size=9.5, color=MUTED)
            if s.get("detail"):
                d = tc.add_paragraph()
                fmt(d, after=0, line=1.15)
                add_run(d, s["detail"])

        if ns.get("closing"):
            p = self.doc.add_paragraph()
            fmt(p, before=18, after=0)
            add_run(p, ns["closing"], italic=True, color=MUTED, size=11.5)

    # -- build ------------------------------------------------------------------
    def build(self, out_path):
        self._header_footer()
        self.cover()
        self.executive_summary()
        self.whats_working()
        self.opportunities()
        self.recommendations()
        self.next_steps()
        self.doc.save(out_path)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("content", help="Path to content JSON")
    ap.add_argument("output", nargs="?", help="Output .docx path")
    ap.add_argument("--check-only", action="store_true", help="Validate content without building")
    args = ap.parse_args()

    with open(args.content, encoding="utf-8") as f:
        content = json.load(f)

    errors, warnings = validate(content)
    for w in warnings:
        print(f"WARNING: {w}")
    if errors:
        for e in errors:
            print(f"ERROR: {e}")
        sys.exit(1)
    if args.check_only:
        print(f"Content OK ({len(warnings)} warning(s)).")
        return
    if not args.output:
        ap.error("output path is required unless --check-only is given")

    ReportBuilder(content).build(args.output)
    print(f"Saved {args.output} ({len(warnings)} warning(s)).")


if __name__ == "__main__":
    main()
